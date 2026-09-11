from __future__ import annotations

import base64
import dataclasses
import io
import os
import pickle
import stat
import sys
import time
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument

import rag_eval_api.parsers.runner as runner_module
from rag_eval_api.parsers.docx import DocxParser
from rag_eval_api.parsers.errors import (
    MalformedDocumentError,
    ParserLimitExceeded,
    ParserSecurityError,
    UnsupportedDocumentError,
)
from rag_eval_api.parsers.models import CanonicalChunk, ParseResult, ParserLimits
from rag_eval_api.parsers.pdf import PdfParser
from rag_eval_api.parsers.registry import ParserRegistry
from rag_eval_api.parsers.runner import (
    ParserRunner,
    ParserSandboxUnavailable,
    ParserTimeout,
    sandbox_command_available,
)


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("Safety", level=1)
    document.add_paragraph("Keep the source location stable.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    stream = b"BT /F1 12 Tf (Hello PDF) Tj ET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(output)


def test_registry_parses_txt_markdown_and_docx_with_deterministic_locations() -> None:
    registry = ParserRegistry()

    text = registry.parse(
        io.BytesIO(b"Heading\n\nFirst paragraph.\n\nSecond paragraph."),
        filename="guide.txt",
        declared_mime="text/plain",
    )
    markdown = registry.parse(
        io.BytesIO(b"# Heading\n\nFirst paragraph."),
        filename="guide.md",
        declared_mime="text/markdown",
    )
    docx = registry.parse(
        io.BytesIO(_docx_bytes()),
        filename="guide.docx",
        declared_mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert [chunk.content for chunk in text.chunks] == [
        "Heading",
        "First paragraph.",
        "Second paragraph.",
    ]
    assert text.chunks[0].source_location == {"paragraph": 0}
    assert markdown.chunks[0].heading == "Heading"
    assert markdown.chunks[0].source_location["line"] == 1
    assert docx.chunks[0].heading == "Safety"
    assert docx.chunks[1].source_location == {"paragraph": 1}
    assert (
        text.content_hash
        == registry.parse(
            io.BytesIO(b"Heading\n\nFirst paragraph.\n\nSecond paragraph."),
            filename="guide.txt",
            declared_mime="text/plain",
        ).content_hash
    )


def test_registry_parses_pdf_and_reports_page_locations() -> None:
    pdf = _pdf_bytes()
    result = ParserRegistry().parse(
        io.BytesIO(pdf), filename="guide.pdf", declared_mime="application/pdf"
    )

    assert result.page_count == 1
    assert result.chunks[0].source_location == {"page": 1}
    assert "Hello PDF" in result.chunks[0].content


@pytest.mark.parametrize(
    ("filename", "mime", "payload"),
    [
        ("guide.pdf", "text/plain", b"%PDF-1.4"),
        ("guide.exe", "application/octet-stream", b"MZ not supported"),
        ("guide.docx", "application/zip", b"not a zip"),
    ],
)
def test_registry_classifies_unsupported_and_malformed_inputs(
    filename: str, mime: str, payload: bytes
) -> None:
    with pytest.raises((UnsupportedDocumentError, MalformedDocumentError)):
        ParserRegistry().parse(io.BytesIO(payload), filename=filename, declared_mime=mime)


def test_registry_rejects_limits_and_zip_traversal(tmp_path: Path) -> None:
    with pytest.raises(ParserLimitExceeded, match="characters"):
        ParserRegistry().parse(
            io.BytesIO(b"one\n\ntwo"),
            filename="guide.txt",
            declared_mime="text/plain",
            limits=ParserLimits(max_normalized_characters=5),
        )

    malicious = io.BytesIO()
    with zipfile.ZipFile(malicious, "w") as archive:
        archive.writestr("../evil.txt", "bad")
    with pytest.raises(MalformedDocumentError, match="ZIP"):
        ParserRegistry().parse(
            io.BytesIO(malicious.getvalue()),
            filename="guide.docx",
            declared_mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    symlinked = io.BytesIO()
    info = zipfile.ZipInfo("word/link.xml")
    info.create_system = 3
    info.external_attr = 0o120777 << 16
    with zipfile.ZipFile(symlinked, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr(info, "../secret")
    with pytest.raises(ParserSecurityError, match="symlink"):
        ParserRegistry().parse(
            io.BytesIO(symlinked.getvalue()),
            filename="guide.docx",
            declared_mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    assert ParserSecurityError.code == "security_violation"


def test_registry_treats_filename_as_metadata_only() -> None:
    result = ParserRegistry().parse(
        io.BytesIO(b"safe content"),
        filename="../../private/secrets.txt",
        declared_mime="text/plain",
    )

    assert result.chunks[0].content == "safe content"


def test_registry_accepts_empty_text_but_rejects_empty_or_truncated_pdf() -> None:
    empty = ParserRegistry().parse(
        io.BytesIO(b""), filename="empty.txt", declared_mime="text/plain"
    )
    assert empty.chunks == ()

    with pytest.raises(UnsupportedDocumentError):
        ParserRegistry().parse(
            io.BytesIO(b""), filename="empty.pdf", declared_mime="application/pdf"
        )


@pytest.mark.parametrize("length", [150_000, 200_000])
def test_text_parser_accepts_long_single_lines_up_to_character_limit(length: int) -> None:
    result = ParserRegistry(isolated=False).parse(
        io.BytesIO(b"x" * length),
        filename="long.txt",
        declared_mime="text/plain",
        limits=ParserLimits(max_normalized_characters=200_000),
    )

    assert sum(chunk.character_count for chunk in result.chunks) == length


def test_text_parser_rejects_single_line_over_character_limit() -> None:
    with pytest.raises(ParserLimitExceeded, match="characters"):
        ParserRegistry(isolated=False).parse(
            io.BytesIO(b"x" * 200_001),
            filename="long.txt",
            declared_mime="text/plain",
            limits=ParserLimits(max_normalized_characters=200_000),
        )


def test_text_parser_handles_multibyte_utf8_split_at_decode_chunk_boundary() -> None:
    content = ("界" * 65_535) + "🙂" + ("界" * 10)
    result = ParserRegistry(isolated=False).parse(
        io.BytesIO(content.encode("utf-8")),
        filename="multibyte.txt",
        declared_mime="text/plain",
        limits=ParserLimits(max_normalized_characters=65_546),
    )

    assert "🙂" in "".join(chunk.content for chunk in result.chunks)
    assert sum(chunk.character_count for chunk in result.chunks) == len(content)
    with pytest.raises(MalformedDocumentError):
        ParserRegistry().parse(
            io.BytesIO(b"%PDF-1.7"), filename="truncated.pdf", declared_mime="application/pdf"
        )


def test_docx_compression_ratio_is_bounded_before_document_parsing() -> None:
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types />", compress_type=zipfile.ZIP_STORED)
        archive.writestr("word/document.xml", b"A" * 50_000)

    with pytest.raises(ParserLimitExceeded, match="compression ratio"):
        ParserRegistry().parse(
            io.BytesIO(bomb.getvalue()),
            filename="bomb.docx",
            declared_mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            limits=ParserLimits(max_docx_compression_ratio=2),
        )


def test_docx_xml_depth_is_bounded_before_document_parsing() -> None:
    deep_xml = ("<a>" * 5) + ("</a>" * 5)
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", deep_xml)

    with pytest.raises(ParserLimitExceeded, match="XML depth"):
        ParserRegistry().parse(
            io.BytesIO(payload.getvalue()),
            filename="deep.docx",
            declared_mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            limits=ParserLimits(max_docx_xml_depth=3),
        )


@pytest.mark.parametrize(
    ("field", "message"),
    [
        ("max_pdf_objects", "object"),
        ("max_pdf_decoded_stream_bytes", "decoded"),
        ("max_pdf_recursion_depth", "recursion"),
    ],
)
def test_pdf_resource_limits_are_enforced_in_isolated_runner(field: str, message: str) -> None:
    with pytest.raises(ParserLimitExceeded, match=message):
        ParserRegistry().parse(
            io.BytesIO(_pdf_bytes()),
            filename="bounded.pdf",
            declared_mime="application/pdf",
            limits=ParserLimits(**{field: 1}),
        )


@pytest.mark.parametrize(
    "field",
    [
        "max_input_bytes",
        "max_pages",
        "max_paragraphs",
        "max_normalized_characters",
        "max_chunk_characters",
        "max_pdf_objects",
        "max_pdf_decoded_stream_bytes",
        "max_pdf_recursion_depth",
        "max_docx_zip_entries",
        "max_docx_uncompressed_bytes",
        "max_docx_compression_ratio",
        "max_docx_xml_depth",
    ],
)
def test_parser_limits_reject_values_above_safe_upper_bounds(field: str) -> None:
    upper_bound = ParserLimits.safe_upper_bounds()[field]
    with pytest.raises(ValueError, match="safe upper bound"):
        ParserLimits(**{field: upper_bound + 1})


def test_parser_runner_enforces_hard_timeout_before_starting_work() -> None:
    with pytest.raises(ParserTimeout):
        ParserRegistry().parse(
            io.BytesIO(_pdf_bytes()),
            filename="guide.pdf",
            declared_mime="application/pdf",
            timeout_seconds=0,
        )


def test_parser_runner_rejects_input_above_parser_limit_before_spawning() -> None:
    with pytest.raises(ParserLimitExceeded, match="input bytes"):
        ParserRunner().parse_bytes(
            b"x" * 11,
            suffix=".pdf",
            filename="oversized.pdf",
            declared_mime="application/pdf",
            limits=ParserLimits(max_input_bytes=10),
        )


def _write_executable(path: Path, content: str) -> None:
    path.write_text(content)
    path.chmod(0o700)


def test_parser_runner_terminates_sleeping_sandbox_process_without_orphan(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        "rag_eval_api.parsers.runner.sandbox_command_available",
        lambda _executable, _args: True,
    )
    pid_file = tmp_path / "sleep.pid"
    wrapper = tmp_path / "bwrap"
    wrapper.write_text(
        "#!" + sys.executable + "\n"
        "import pathlib, subprocess\n"
        f"pid_file = pathlib.Path({str(pid_file)!r})\n"
        "child = subprocess.Popen(['sleep', '30'])\n"
        "pid_file.write_text(str(child.pid))\n"
        "child.wait()\n"
    )
    wrapper.chmod(0o700)
    runner = ParserRunner(
        sandbox_executable=str(wrapper),
        sandbox_args=("--die-with-parent", "--unshare-net"),
    )

    with pytest.raises(ParserTimeout):
        runner.parse_bytes(
            _pdf_bytes(),
            suffix=".pdf",
            filename="timeout.pdf",
            declared_mime="application/pdf",
            limits=ParserLimits(),
            timeout_seconds=2.0,
        )

    deadline = time.monotonic() + 2
    while not pid_file.exists() and time.monotonic() < deadline:
        time.sleep(0.02)
    assert pid_file.exists()
    child_pid = int(pid_file.read_text())
    for _ in range(40):
        try:
            os.kill(child_pid, 0)
        except ProcessLookupError:
            break
        time.sleep(0.05)
    else:
        pytest.fail("timed-out parser descendant was left running")


@pytest.mark.parametrize("exit_code", [0, 17])
def test_parser_runner_maps_eof_or_crashed_child_to_sandbox_error(
    tmp_path: Path, exit_code: int, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        "rag_eval_api.parsers.runner.sandbox_command_available",
        lambda _executable, _args: True,
    )
    wrapper = tmp_path / "bwrap"
    _write_executable(wrapper, f"#!/bin/sh\nexit {exit_code}\n")

    with pytest.raises(ParserSandboxUnavailable):
        ParserRunner(
            sandbox_executable=str(wrapper),
            sandbox_args=("--die-with-parent", "--unshare-net"),
        ).parse_bytes(
            _pdf_bytes(),
            suffix=".pdf",
            filename="crashed.pdf",
            declared_mime="application/pdf",
            limits=ParserLimits(),
        )


def test_parser_runner_requires_os_sandbox_for_strict_mode() -> None:
    with pytest.raises(ParserSandboxUnavailable):
        ParserRunner(require_resource_limits=True).parse_bytes(
            _pdf_bytes(),
            suffix=".pdf",
            filename="strict.pdf",
            declared_mime="application/pdf",
            limits=ParserLimits(),
        )


def test_parser_runner_rejects_pickle_payload_without_executing_it(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        "rag_eval_api.parsers.runner.sandbox_command_available",
        lambda _executable, _args: True,
    )
    marker = tmp_path / "executed"

    class MaliciousPayload:
        def __reduce__(self) -> tuple[object, tuple[str]]:
            return (os.system, (f"touch {marker}",))

    encoded = base64.b64encode(pickle.dumps(MaliciousPayload())).decode("ascii")
    wrapper = tmp_path / "bwrap"
    wrapper.write_text(
        "#!" + sys.executable + "\n"
        "import base64, sys\n"
        f"sys.stdout.buffer.write(base64.b64decode({encoded!r}))\n"
    )
    wrapper.chmod(0o700)

    with pytest.raises(ParserSandboxUnavailable):
        ParserRunner(
            sandbox_executable=str(wrapper),
            sandbox_args=("--die-with-parent", "--unshare-net"),
        ).parse_bytes(
            _pdf_bytes(),
            suffix=".pdf",
            filename="malicious.pdf",
            declared_mime="application/pdf",
            limits=ParserLimits(),
        )
    assert not marker.exists()


def test_parser_runner_rejects_sandbox_commands_without_required_flags(tmp_path: Path) -> None:
    bwrap = tmp_path / "bwrap"
    bwrap.write_text("#!/bin/sh\nexit 0\n")
    bwrap.chmod(0o700)
    assert not sandbox_command_available(str(bwrap), ("--die-with-parent",))
    assert not sandbox_command_available(str(bwrap), ("--die-with-parent", "--unshare-net"))
    env = tmp_path / "env"
    env.write_text("#!/bin/sh\nexit 0\n")
    env.chmod(0o700)
    assert not sandbox_command_available(str(env), ("--unshare-net", "--die-with-parent"))


def test_sandbox_template_requires_exact_order_and_rejects_extra_args(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bwrap = tmp_path / "bwrap"
    _write_executable(bwrap, "#!/bin/sh\nexit 0\n")
    monkeypatch.setattr(
        runner_module,
        "_SANDBOX_PATHS",
        {"bwrap": frozenset({str(bwrap)})},
    )
    monkeypatch.setattr(
        runner_module.os,
        "stat",
        lambda *_args, **_kwargs: SimpleNamespace(st_uid=0, st_mode=0o100755),
    )

    monkeypatch.setattr(runner_module, "_probe_sandbox_template", lambda *_args: True)
    exact = runner_module._SANDBOX_TEMPLATES["bwrap"]
    assert sandbox_command_available(str(bwrap), exact)
    assert not sandbox_command_available(str(bwrap), (*exact, "--ro-bind", "/", "/"))
    assert not sandbox_command_available(str(bwrap), tuple(reversed(exact)))


def test_unshare_is_rejected_without_a_filesystem_boundary(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    unshare = tmp_path / "unshare"
    _write_executable(unshare, "#!/bin/sh\nexit 0\n")
    monkeypatch.setattr(
        runner_module,
        "_SANDBOX_PATHS",
        {"unshare": frozenset({str(unshare)})},
    )
    monkeypatch.setattr(
        runner_module.os,
        "stat",
        lambda *_args, **_kwargs: SimpleNamespace(st_uid=0, st_mode=0o100755),
    )

    assert not sandbox_command_available(
        str(unshare),
        (
            "--user",
            "--map-root-user",
            "--mount",
            "--uts",
            "--ipc",
            "--net",
            "--pid",
            "--fork",
            "--kill-child",
        ),
    )


def test_parser_runner_places_parser_argv_after_sandbox_separator(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bwrap = tmp_path / "bwrap"
    _write_executable(bwrap, "#!/bin/sh\nexit 0\n")
    monkeypatch.setattr(
        runner_module,
        "sandbox_command_available",
        lambda _executable, _args: True,
    )
    monkeypatch.setattr(runner_module, "_sandbox_runtime_bind_args", lambda: [])
    command = ParserRunner(
        sandbox_executable=str(bwrap),
        sandbox_args=("--unshare-net", "--die-with-parent", "--new-session"),
    )._command()
    assert command[:5] == [
        str(bwrap),
        "--unshare-net",
        "--die-with-parent",
        "--new-session",
        "--",
    ]
    assert command[5:] == [
        sys.executable,
        "-m",
        "rag_eval_api.parsers.runner",
        "--child",
        "--strict",
    ]


def test_parser_runner_uses_minimal_environment_and_private_working_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    result = ParseResult(
        parser_version="test-v1",
        content_hash="0" * 64,
        byte_size=1,
        page_count=0,
        paragraph_count=0,
    )
    output = runner_module._json_bytes(
        {
            "kind": "ok",
            "result": runner_module._result_envelope(result).model_dump(mode="json"),
        }
    )
    captured: dict[str, object] = {}

    class FakeProcess:
        pid = os.getpid()
        returncode = 0
        stdin = None

        def poll(self) -> int:
            return 0

    def fake_popen(command: list[str], **kwargs: object) -> FakeProcess:
        captured["command"] = command
        captured.update(kwargs)
        cwd = Path(str(kwargs["cwd"]))
        assert cwd.is_dir()
        assert stat.S_IMODE(cwd.stat().st_mode) == 0o700
        return FakeProcess()

    monkeypatch.setenv("DATABASE_URL", "postgresql://secret")
    monkeypatch.setenv("PROVIDER_API_KEY", "secret")
    monkeypatch.setattr(runner_module.subprocess, "Popen", fake_popen)
    monkeypatch.setattr(
        ParserRunner,
        "_communicate_bounded",
        lambda self, process, request, timeout, output_limit: (output, b""),
    )

    ParserRunner().parse_bytes(
        b"safe",
        suffix=".pdf",
        filename="safe.pdf",
        declared_mime="application/pdf",
        limits=ParserLimits(),
    )

    child_env = captured["env"]
    assert isinstance(child_env, dict)
    assert "DATABASE_URL" not in child_env
    assert "PROVIDER_API_KEY" not in child_env
    assert child_env["PYTHONPATH"].split(os.pathsep)[0] == str(
        Path(runner_module.__file__).resolve().parents[2]
    )
    assert child_env["TMPDIR"] == captured["cwd"]
    assert not Path(str(captured["cwd"])).exists()


def test_sandbox_runtime_binds_reject_broad_prefixes(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(runner_module.sys, "executable", "/")
    assert runner_module._sandbox_runtime_bind_args() == []


def test_runtime_python_paths_reject_untrusted_sysconfig_locations(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        runner_module.sysconfig,
        "get_path",
        lambda _key: "/tmp/python-runtime/site-packages",
    )
    assert runner_module._runtime_python_paths() == set()


@pytest.mark.skipif(
    sys.platform != "linux" or not Path("/usr/bin/bwrap").is_file(),
    reason="Linux bubblewrap is required for the real sandbox capability check",
)
def test_real_bwrap_profile_passes_capability_probe() -> None:
    if os.geteuid() == 0:
        pytest.skip("the real profile must be exercised by a non-root process")
    assert runner_module.sandbox_command_available(
        "/usr/bin/bwrap", runner_module._SANDBOX_TEMPLATES["bwrap"]
    )


def test_child_main_applies_fixed_limits_before_reading_protocol(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    events: list[str] = []

    def startup_limits() -> None:
        events.append("startup-limits")

    def read_protocol(*args: object, **kwargs: object) -> bytes:
        del args, kwargs
        events.append("read-protocol")
        raise ParserLimitExceeded("protocol size")

    monkeypatch.setattr(runner_module, "_configure_startup_limits", startup_limits, raising=False)
    monkeypatch.setattr(runner_module, "_read_bounded", read_protocol)
    stdout = io.BytesIO()
    monkeypatch.setattr(runner_module.sys, "stdout", SimpleNamespace(buffer=stdout))
    monkeypatch.setattr(runner_module.sys, "stdin", SimpleNamespace(buffer=io.BytesIO()))

    runner_module._child_main()

    assert events == ["startup-limits", "read-protocol"]
    envelope = runner_module._strict_json_loads(stdout.getvalue())
    assert isinstance(envelope, dict)
    assert envelope["code"] == "size_exceeded"


def test_parser_input_safe_upper_bound_matches_parent_protocol_budget() -> None:
    assert ParserLimits.safe_upper_bounds()["max_input_bytes"] == 64 * 1024 * 1024


def test_parser_runner_output_limit_is_the_explicit_serialized_budget() -> None:
    assert runner_module._child_output_limit(1 * 1024 * 1024) == 1 * 1024 * 1024
    assert runner_module._child_output_limit(64 * 1024 * 1024) == 64 * 1024 * 1024


@pytest.mark.parametrize("value", [0, 1024 * 1024 - 1, 256 * 1024 * 1024 + 1, True])
def test_parser_runner_output_budget_is_explicit_and_bounded(value: object) -> None:
    with pytest.raises(ValueError, match="output"):
        runner_module._child_output_limit(value)  # type: ignore[arg-type]


def test_default_parser_output_budget_allows_5000_legal_chunks() -> None:
    result = ParseResult(
        parser_version="test-v1",
        content_hash="0" * 64,
        byte_size=10_000_000,
        page_count=0,
        paragraph_count=5000,
        chunks=tuple(
            CanonicalChunk.create(
                ordinal=index,
                content="x" * 2000,
                source_location={"paragraph": index},
            )
            for index in range(5000)
        ),
    )
    payload = runner_module._json_bytes(
        {
            "kind": "ok",
            "result": runner_module._result_envelope(result).model_dump(mode="json"),
        }
    )
    assert len(payload) < 64 * 1024 * 1024
    assert len(runner_module._parse_result_or_error(payload).chunks) == 5000


def test_parser_runner_rejects_request_over_protocol_budget_before_spawn(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(runner_module, "_MAX_CHILD_REQUEST_BYTES", 4)
    with pytest.raises(ParserLimitExceeded, match="configured input budget"):
        ParserRunner().parse_bytes(
            _pdf_bytes(),
            suffix=".pdf",
            filename="oversized-request.pdf",
            declared_mime="application/pdf",
            limits=ParserLimits(),
        )


def test_child_sandbox_applies_cpu_and_address_space_limits(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import resource

    calls: list[tuple[int, tuple[int, int]]] = []
    monkeypatch.setattr(
        resource,
        "setrlimit",
        lambda resource_kind, values: calls.append((resource_kind, values)),
    )
    monkeypatch.setattr(runner_module, "_resource_limits_supported", lambda: True)

    runner_module._configure_child_sandbox(
        ParserLimits(max_input_bytes=10 * 1024 * 1024),
        2.0,
        sandbox_enabled=True,
    )

    assert {resource_kind for resource_kind, _ in calls} == {
        resource.RLIMIT_CPU,
        resource.RLIMIT_AS,
    }


def test_child_result_rejects_oversized_base64_before_decoding(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(runner_module, "_configure_child_sandbox", lambda *args, **kwargs: None)
    request = runner_module._RequestEnvelope.model_validate(
        {
            "data_b64": base64.b64encode(b"x" * 11).decode("ascii"),
            "suffix": ".pdf",
            "filename": "oversized.pdf",
            "declared_mime": "application/pdf",
            "limits": dataclasses.asdict(ParserLimits(max_input_bytes=10)),
            "timeout_seconds": 1.0,
            "sandbox_enabled": False,
        }
    )

    result = runner_module._child_result(request)
    assert isinstance(result, runner_module._ErrorEnvelope)
    assert result.code == "size_exceeded"


def test_child_protocol_read_is_bounded() -> None:
    with pytest.raises(ParserLimitExceeded, match="protocol size"):
        runner_module._read_bounded(io.BytesIO(b"12345"), max_bytes=4)


@pytest.mark.parametrize(
    "payload",
    [
        b'{"kind":"ok","kind":"error"}',
        b'{"kind":"ok","value":NaN}',
        b'{"kind":"ok",',
    ],
)
def test_parser_protocol_rejects_duplicate_nan_and_invalid_json(payload: bytes) -> None:
    with pytest.raises(ParserSandboxUnavailable):
        runner_module._parse_result_or_error(payload)


def test_child_main_returns_stable_error_for_bounded_oversized_request(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(runner_module, "_MAX_CHILD_REQUEST_BYTES", 4)
    stdin = SimpleNamespace(buffer=io.BytesIO(b"12345"))
    stdout = io.BytesIO()
    monkeypatch.setattr(runner_module.sys, "stdin", stdin)
    monkeypatch.setattr(runner_module.sys, "stdout", SimpleNamespace(buffer=stdout))

    runner_module._child_main()

    envelope = runner_module._strict_json_loads(stdout.getvalue())
    assert isinstance(envelope, dict)
    assert envelope["kind"] == "error"
    assert envelope["code"] == "size_exceeded"


def test_registry_keeps_pdf_and_docx_on_runner_even_when_text_is_unisolated() -> None:
    class StubRunner:
        def __init__(self) -> None:
            self.calls: list[str] = []

        def parse_bytes(self, data: bytes, **kwargs: object) -> object:
            del data
            self.calls.append(str(kwargs["suffix"]))
            return object()

    runner = StubRunner()
    registry = ParserRegistry(isolated=False, runner=runner)  # type: ignore[arg-type]
    registry.parse(io.BytesIO(_pdf_bytes()), filename="safe.pdf", declared_mime="application/pdf")
    assert runner.calls == [".pdf"]


@pytest.mark.parametrize("stream", ["stdout", "stderr"])
def test_parser_runner_caps_child_output_and_kills_process_group(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, stream: str
) -> None:
    monkeypatch.setattr(
        "rag_eval_api.parsers.runner.sandbox_command_available",
        lambda _executable, _args: True,
    )
    wrapper = tmp_path / "bwrap"
    wrapper.write_text(
        "#!" + sys.executable + "\n"
        "import sys\n"
        f"sys.{stream}.write('x' * (2 * 1024 * 1024))\n"
        f"sys.{stream}.flush()\n"
    )
    wrapper.chmod(0o700)

    with pytest.raises(ParserSandboxUnavailable, match="output"):
        ParserRunner(
            max_output_bytes=1 * 1024 * 1024,
            sandbox_executable=str(wrapper),
            sandbox_args=("--die-with-parent", "--unshare-net"),
        ).parse_bytes(
            _pdf_bytes(),
            suffix=".pdf",
            filename="output-limit.pdf",
            declared_mime="application/pdf",
            limits=ParserLimits(),
        )


def test_pdf_and_docx_parser_timeout_is_not_downgraded_to_parse_failed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_clock = iter((0.0, 2.0))
    monkeypatch.setattr("rag_eval_api.parsers.pdf.time.monotonic", lambda: next(pdf_clock))
    with pytest.raises(ParserTimeout) as pdf_error:
        PdfParser().parse(
            _pdf_bytes(),
            filename="timeout.pdf",
            declared_mime="application/pdf",
            limits=ParserLimits(max_pdf_wall_clock_seconds=1),
        )
    assert pdf_error.value.code == "parse_timeout"

    docx_clock = iter((0.0, 2.0))
    monkeypatch.setattr("rag_eval_api.parsers.docx.time.monotonic", lambda: next(docx_clock))
    with pytest.raises(ParserTimeout) as docx_error:
        DocxParser().parse(
            _docx_bytes(),
            filename="timeout.docx",
            declared_mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            limits=ParserLimits(max_pdf_wall_clock_seconds=1),
        )
    assert docx_error.value.code == "parse_timeout"
